#!/usr/bin/env python3
"""사용기 분석 추출 정확도 평가 — 채점과 보고서(기획 7장 M1~M13).

DB 를 쓰지 않는다. `data/eval_extraction/model_runs/*.json`(모델 3회 결과)와 정답을 비교한다.
정답이 아직 없으면 정답이 필요 없는 지표(M9 환각, M13 재현성, 오류, 항목 수)만 낸다.

    python scripts/eval_extraction_score.py                                   # 정답 없이: 정답이 필요 없는 지표
    python scripts/eval_extraction_score.py --gold data/eval_extraction/gold.json
    python scripts/eval_extraction_score.py --gold-xlsx data/eval_extraction/label_sheets_A_done.xlsx   # 한 라벨러 시트로 임시 채점
    python scripts/eval_extraction_score.py --gold-xlsx A.xlsx --gold-xlsx B.xlsx   # 두 시트의 일치도(문장 단위)도 낸다
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.engine import peripheral_post_eval_metrics as metrics    # noqa: E402

DIR = ROOT / "data" / "eval_extraction"

YES = {"예", "y", "yes", "true", "1", "o"}


def _yes(v) -> bool:
    return str(v or "").strip().lower() in YES


def _csv(v) -> list[str]:
    return [x.strip() for x in str(v or "").replace("，", ",").split(",") if x.strip()]


def load_gold_xlsx(path: Path, key: dict[str, dict], multi_ids: set[str]) -> dict[str, dict]:
    """라벨러가 채운 시트(글마다 P01~ 시트) → 정답 JSON. 비어 있는 시트는 건너뛴다."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    gold: dict[str, dict] = {}
    for name in wb.sheetnames:
        if name not in key:
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        fields: dict[str, object] = {}
        items: list[dict] = []
        header_at = None
        for i, row in enumerate(rows):
            label = str(row[0] or "")
            if label.startswith("문장(그대로 복사)"):
                header_at = i
                break
            for tag in ("사용기인가", "아니라면 사유", "본문에 지시문", "손 길이", "손 크기 단계", "용도", "상황", "사용 기간"):
                if label.startswith(tag):
                    fields[tag] = row[1]
        if header_at is None or not fields.get("사용기인가"):
            continue
        for row in rows[header_at + 1:]:
            if not row[0]:
                continue
            items.append({
                "sentence": str(row[0]).strip(),
                "product_index": int(row[1]) if row[1] not in (None, "") else 0,
                "experienced": _yes(row[2]), "comparison": _yes(row[3]), "spec_statement": _yes(row[4]),
                "attribute": str(row[5] or "").strip() or "other",
                "direction": str(row[6]).strip() if row[6] else None,
                "side": str(row[7]).strip() if row[7] else None,
                "evaluated_in": _csv(row[8]),
            })
        ext = key[name]["post_id_ext"]
        hand_cm = fields.get("손 길이")
        gold[ext] = {
            "is_review": _yes(fields.get("사용기인가")),
            "reason": str(fields.get("아니라면 사유") or "").strip() or None,
            "has_injection": _yes(fields.get("본문에 지시문")),
            "hand_cm": float(hand_cm) if hand_cm not in (None, "") else None,
            "hand_size": str(fields.get("손 크기 단계") or "").strip() or None,
            "purposes": _csv(fields.get("용도")), "situations": _csv(fields.get("상황")),
            "multi_product": ext in multi_ids, "items": items,
        }
    return gold


def agreement(a: dict, b: dict) -> dict:
    """두 라벨러의 문장 단위 일치도: 같은 (문장·제품) 키를 둘 다 적은 비율, 속성·방향 일치율. (Cohen's κ 는 문장 후보 집합이 없어 근사)"""
    keys_a = {(metrics.squash(i["sentence"]), i["product_index"]): i for g in a.values() for i in g["items"]}
    keys_b = {(metrics.squash(i["sentence"]), i["product_index"]): i for g in b.values() for i in g["items"]}
    both = set(keys_a) & set(keys_b)
    attr = sum(1 for k in both if keys_a[k]["attribute"] == keys_b[k]["attribute"])
    value = sum(1 for k in both if keys_a[k]["attribute"] == keys_b[k]["attribute"]
                and (keys_a[k]["direction"] or keys_a[k]["side"]) == (keys_b[k]["direction"] or keys_b[k]["side"]))
    union = set(keys_a) | set(keys_b)
    return {"sentence_overlap": len(both) / len(union) if union else None, "attribute_match": attr / len(both) if both else None,
            "value_match_given_attribute": value / attr if attr else None, "n_both": len(both), "n_union": len(union)}


def fmt(v) -> str:
    return "—" if v is None else (f"{v:.0%}" if isinstance(v, float) and v <= 1 else str(v))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--gold", type=Path)
    ap.add_argument("--gold-xlsx", type=Path, action="append")
    ap.add_argument("--out", type=Path, default=DIR / "report.md")
    args = ap.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    manifest = json.loads((DIR / "manifest.json").read_text(encoding="utf-8"))
    kinds = {p["post_id_ext"]: p["kind"] for p in manifest["posts"]}
    multi = {p["post_id_ext"] for p in manifest["posts"] if p.get("multi_product")}
    records = {}
    for path in sorted((DIR / "model_runs").glob("*.json")):
        rec = json.loads(path.read_text(encoding="utf-8"))
        records[rec["post_id_ext"]] = rec

    lines = [f"# 사용기 분석 추출 정확도 평가 보고서(초안)", "",
             f"- 말뭉치: {len(manifest['posts'])}편 · 모델 결과 {len(records)}편 · 동결: git `{manifest['freeze']['git_head'][:10]}`"
             f"{' (작업 트리에 미커밋 변경 있음)' if manifest['freeze']['working_tree_dirty_for_frozen_files'] else ''}",
             f"- 동결 파일 해시: {manifest['freeze']['file_sha256_16']}", ""]
    free = metrics.gold_free(records)
    lines += ["## 정답이 필요 없는 지표", "", "| 항목 | 값 |", "|---|---|",
              f"| 글 · 호출 · 오류 | {free['posts']} · {free['runs']} · {free['run_errors']} |",
              f"| 원래 추출(raw) 항목 / 최종(confirmed) 항목 | {free['raw_items']} / {free['confirmed_items']} |",
              f"| M9 본문에 없는 문장(raw) | {free['M9_hallucinated_sentences']} |",
              f"| M13 재현성(같은 글 회차 간 최종 항목 Jaccard 평균) | {fmt(free['M13_reproducibility_jaccard'])} |", ""]

    gold: dict[str, dict] = {}
    if args.gold:
        gold = json.loads(args.gold.read_text(encoding="utf-8"))
    xlsx_gold: list[dict] = []
    if args.gold_xlsx:
        key = json.loads((DIR / "label_key.json").read_text(encoding="utf-8"))
        xlsx_gold = [load_gold_xlsx(p, key, multi) for p in args.gold_xlsx]
        if not gold:
            gold = xlsx_gold[0]
    if len(xlsx_gold) == 2:
        ag = agreement(xlsx_gold[0], xlsx_gold[1])
        lines += ["## 라벨러 일치도(근사)", "", f"- {ag}", ""]
    if gold:
        result = metrics.score(records, gold, kinds)
        lines += [f"## 정답 비교 (정답 {len(gold)}편)", "", "| 지표 | " + " | ".join(result) + " |",
                  "|---|" + "---|" * len(result)]
        for name in ("M1", "M2", "M3", "M4", "M5", "M6", "M7", "M8_spec", "M8_comparison", "M8_guess", "M10_extra_values",
                     "M11", "M12_held_rate"):
            lines.append(f"| {name} | " + " | ".join(fmt(result[b].get(name)) for b in result) + " |")
        lines.append("")
    else:
        lines += ["## 정답 비교", "", "정답(라벨 시트)이 아직 없어 M1~M8·M10~M12 는 계산하지 않았다.", ""]
    text = "\n".join(lines) + "\n"
    args.out.write_text(text, encoding="utf-8")
    print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
