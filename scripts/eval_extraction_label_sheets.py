#!/usr/bin/env python3
"""사용기 분석 추출 정확도 평가 — 라벨링 시트(xlsx) 생성(기획 4장·8장).

글마다 시트 한 장: 연결 제품, 분석 입력 본문(서식을 걷어낸 글자 본문), 글 단위 칸, 작성자 조건 칸, 문장 라벨 표.
라벨러가 편향되지 않도록 **글 번호를 가리고(P01~) 순서를 섞는다.** 열쇠(P번호 → 글 번호)는 `label_key.json` 에 따로 두고 라벨러에게 주지 않는다.
모델 결과는 어떤 칸에도 넣지 않는다. 라벨러 2명이 같은 시트를 각자 채운다(`--labelers A,B`).

    python scripts/eval_extraction_label_sheets.py --labelers A,B
"""
from __future__ import annotations

import argparse
import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.db import get_conn                                                    # noqa: E402
from src.engine.peripheral_rules import (feel_attributes, load_peripheral_rules,   # noqa: E402
                                         purpose_codes, situations)
from src.repo import peripheral_post_source as source                          # noqa: E402

OUT_DIR = ROOT / "data" / "eval_extraction"
REASONS = ["question", "chatter", "ad", "unboxing", "planned", "store_demo", "hearsay", "sale"]
FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="D9E1F2")
SENTENCE_ROWS = 30


def _code_lines(rules) -> list[tuple[str, str, str]]:
    out = []
    for kind, label in (("mouse", "마우스"), ("keyboard", "키보드"), ("monitor", "모니터"), ("speaker", "스피커")):
        for code, e in feel_attributes(kind, rules).items():
            if e["type"] == "rating":
                out.append((label, code, f"{e['label']} — 방향 good/bad/mixed"))
            else:
                sides = ", ".join(f"{s}={lab}" for s, lab in e["sides"].items())
                out.append((label, code, f"{e['label']} — 쪽 {sides}"))
    return out


def build(labeler: str, order: list[dict], rules, bodies: dict[str, dict]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "안내"
    guide = [
        f"라벨러 {labeler} — 사용기 분석 추출 정확도 평가 정답 라벨링",
        "",
        "1. 글마다 시트가 한 장씩 있다(P01~). 본문을 읽고 노란 칸을 채운다. 모델이 어떻게 읽었는지는 보지 않는다.",
        "2. 글 단위: 사용기인가(예/아니오). 아니라면 사유 코드, 본문에 지시문(\"AI는 ~하라\")이 있는지.",
        "3. 작성자 조건: 글에 밝힌 것만. 손 길이(cm)는 cm로 적은 경우만, 손 크기 단계는 \"작은 편/보통/큰 편\"이라고만 적은 경우만.",
        "   용도·상황은 글쓴이가 \"나는 이걸 무엇에/어디서 쓴다\"고 스스로 밝힌 것만(기능 평가 중 스친 장면은 제외). 없으면 비운다.",
        "4. 문장 라벨: 본문에 있는 체감 문장을 한 글자도 바꾸지 말고 복사한다. 한 문장에 속성이 여럿이면 행을 나눈다.",
        "   평가 속성은 방향(good/bad/mixed), 성향 속성은 쪽을 채운다(둘을 같이 채우지 않는다). 맞는 속성이 없으면 other.",
        "   스펙 옮김·비교·짐작·전해 들은 말도 라벨은 하되 해당 칸을 \"예\"로 표시한다(모델이 걸러내는지 보려는 것).",
        "5. 애매하면 other 로 두고 '라벨러 메모'에 적는다. 억지로 속성에 넣지 않는다.",
        "6. 다 채운 파일은 라벨러가 서로 보지 않고 제출한다. 불일치는 둘이 합의해 최종 정답으로 확정한다.",
        "",
        "속성 코드표는 '코드표' 시트.",
    ]
    for i, line in enumerate(guide, start=1):
        ws.cell(i, 1, line)
    ws.column_dimensions["A"].width = 120
    ws["A1"].font = Font(bold=True, size=13)

    ct = wb.create_sheet("코드표")
    ct.append(["종류", "속성 코드", "설명"])
    for row in _code_lines(rules):
        ct.append(list(row))
    ct.append([])
    ct.append(["용도(purposes)", ", ".join(purpose_codes(rules))])
    ct.append(["상황(situations)", ", ".join(f"{c}={e['label']}" for c, e in situations(rules).items())])
    ct.append(["비사용기 사유", ", ".join(REASONS)])
    for c in "ABC":
        ct.column_dimensions[c].width = 28 if c != "C" else 70

    attr_codes = sorted({code for _, code, _ in _code_lines(rules)} | {"other"})
    sides = sorted({s for kind in ("mouse", "keyboard", "monitor", "speaker")
                    for e in feel_attributes(kind, rules).values() if e["type"] == "taste" for s in e["sides"]})
    def _dv(options: str) -> DataValidation:
        return DataValidation(type="list", formula1='"' + options + '"', allow_blank=True)

    attr_options = ",".join(attr_codes)
    side_options = ",".join(sides)

    index = wb.create_sheet("목록")
    index.append(["시트", "연결 제품 수", "본문 글자 수", "완료(O)"])
    for entry in order:
        sheet_name = entry["label"]
        info = bodies[entry["post_id"]]
        index.append([sheet_name, len(info["products"]), len(info["body"]), ""])
        s = wb.create_sheet(sheet_name)
        yn, reason_dv, hand_dv = _dv("예,아니오"), _dv(",".join(REASONS)), _dv("small,medium,large")
        dir_dv, attr_dv, side_dv = _dv("good,bad,mixed"), _dv(attr_options), _dv(side_options)
        for dv in (yn, reason_dv, hand_dv, dir_dv, attr_dv, side_dv):
            s.add_data_validation(dv)
        s.append(["글", sheet_name]); s.append([])
        s.append(["연결 제품 번호", "제품", "옵션"])
        for p in info["products"]:
            s.append([p["index"], p["name"], p["option"] or ""])
        s.append([])
        s.append(["본문(분석 입력 글자 본문)"])
        body_row = s.max_row + 1
        s.cell(body_row, 1, info["body"])
        s.merge_cells(start_row=body_row, start_column=1, end_row=body_row, end_column=10)
        s.cell(body_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        s.row_dimensions[body_row].height = min(409, max(60, len(info["body"]) // 3))
        s.append([])
        s.append(["글 단위 라벨"])
        top = s.max_row + 1
        fields = [("사용기인가(예/아니오)", yn), ("아니라면 사유", reason_dv), ("본문에 지시문이 있는가(예/아니오)", yn),
                  ("손 길이(cm) — cm로 적은 경우만", None), ("손 크기 단계(small/medium/large)", hand_dv),
                  ("용도(쉼표로 구분, 코드표)", None), ("상황(쉼표로 구분, 코드표)", None), ("사용 기간(개월) — 밝힌 경우만", None),
                  ("라벨러 메모", None)]
        for i, (label, dv) in enumerate(fields):
            s.cell(top + i, 1, label)
            cell = s.cell(top + i, 2)
            cell.fill = FILL
            if dv is not None:
                dv.add(cell)
        s.append([])
        head_row = s.max_row + 1
        headers = ["문장(그대로 복사)", "제품 번호", "직접 겪은 체감(예/아니오)", "비교 문장(예/아니오)", "스펙 옮김(예/아니오)",
                   "속성", "방향(평가)", "쪽(성향)", "평가한 상황(코드)", "사용 기간(개월)"]
        for c, h in enumerate(headers, start=1):
            cell = s.cell(head_row, c, h)
            cell.font = Font(bold=True)
            cell.fill = HEAD
        for r in range(head_row + 1, head_row + 1 + SENTENCE_ROWS):
            for c in range(1, len(headers) + 1):
                s.cell(r, c).fill = FILL
            for col, dv in ((3, yn), (4, yn), (5, yn), (6, attr_dv), (7, dir_dv), (8, side_dv)):
                dv.add(s.cell(r, col))
        s.column_dimensions["A"].width = 60
        s.column_dimensions["B"].width = 22
        for c in "CDEFGHIJ":
            s.column_dimensions[c].width = 18
    return wb


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--labelers", default="A,B")
    ap.add_argument("--seed", type=int, default=20261014)
    args = ap.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    manifest = json.loads((OUT_DIR / "manifest.json").read_text(encoding="utf-8"))
    posts = [p for p in manifest["posts"] if p.get("post_id")]
    rng = random.Random(args.seed)
    shuffled = posts[:]
    rng.shuffle(shuffled)
    order = [{"label": f"P{i:02d}", "post_id": p["post_id"], "post_id_ext": p["post_id_ext"]} for i, p in enumerate(shuffled, start=1)]
    rules = load_peripheral_rules()
    bodies: dict[str, dict] = {}
    with get_conn() as conn:
        for p in posts:
            post = source.load_post_input(conn, p["post_id"])
            bodies[p["post_id"]] = {
                "body": post.body,
                "products": [{"index": x.index, "name": x.name, "option": x.option_label} for x in post.products],
            }
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "label_key.json").write_text(
        json.dumps({o["label"]: {"post_id": o["post_id"], "post_id_ext": o["post_id_ext"]} for o in order}, ensure_ascii=False, indent=1),
        encoding="utf-8")
    for labeler in [x.strip() for x in args.labelers.split(",") if x.strip()]:
        wb = build(labeler, order, rules, bodies)
        path = OUT_DIR / f"label_sheets_{labeler}.xlsx"
        wb.save(path)
        print("저장:", path)
    print(f"시트 {len(order)}장. 열쇠: {OUT_DIR / 'label_key.json'} (라벨러에게 주지 않는다)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
