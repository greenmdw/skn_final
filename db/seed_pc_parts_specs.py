#!/usr/bin/env python3
"""PC 부품 8종 카탈로그를 DB에 적재 — data/parts_list_modify.xlsx → catalog.*.

시트(CPU/MainBoard/RAM/GPU/SSD/PSU/Case/Cooler)당 한 카테고리. 공통 정보(브랜드·모델·
이미지)는 기존 catalog.product/product_variant, 가격은 기존 catalog.offer/
offer_observation 경로를 그대로 쓴다(seed_catalog.py와 동일 패턴). 카테고리별 스펙은
0015_pc_parts_category_specs.sql 이 만든 catalog.*_spec 테이블에 컬럼으로 저장한다
(EAV 대신 JOIN — 팀 논의 2026-09-18, 소스 데이터가 이미 카테고리별 고정 컬럼 표라서).

멱등(ON CONFLICT) — 몇 번을 다시 돌려도 같은 (brand, model) 행을 갱신할 뿐 중복
생성하지 않는다. 원본 엑셀이 source of truth, DB는 이 스크립트로 언제든 재생성 가능한
산출물이라는 프로젝트 관행을 따른다.

    DATABASE_URL=... python db/seed_pc_parts_specs.py
"""
from __future__ import annotations

import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.config import DATA_DIR  # noqa: E402
from src.db import get_conn  # noqa: E402
from src.repo.product_repo import ProductRepo  # noqa: E402

_XLSX_PATH = DATA_DIR / "parts_list_modify.xlsx"
_SOURCE_NAME = "다나와 수집 데이터 (PC 부품)"
_MERCHANT_PLATFORM = "danawa"
_MERCHANT_SELLER_ID = "danawa-aggregate"
_MERCHANT_NAME = "다나와 가격비교"

# 시트명 -> (product_type 슬러그, 전용 스펙 테이블, {엑셀 헤더: DB 컬럼} 매핑)
_SHEET_SPECS: dict[str, tuple[str, str, dict[str, str]]] = {
    "CPU": ("cpu", "catalog.cpu_spec", {
        "소켓 규격": "socket", "기본 소비전력(W)": "tdp_w", "메모리 규격": "memory_type",
        "내장 그래픽": "integrated_graphics", "쿨러 포함": "cooler_included",
        "제품 라인업": "lineup", "크기": "size_note", "무게": "weight_note",
        "지원 URL": "spec_url", "판매 상태": "sale_status",
        "상태 확인일": "status_checked_at", "비고": "note",
        # 0018: 호환·성능 확장 열(원본에 없으면 NULL)
        "최대 패키지 전력/PL2(PPT)(W)": "max_power_w", "CPU 계열(정규)": "family",
        "전력·계열 출처 URL": "power_family_source_url",
        "성능 점수": "perf_score", "성능 점수 출처 URL": "perf_score_source_url",
    }),
    "MainBoard": ("motherboard", "catalog.mainboard_spec", {
        "소켓 규격": "socket", "칩셋": "chipset", "폼팩터": "form_factor",
        "메모리 규격": "memory_type", "보드 크기(mm)": "board_size_mm",
        "지원 CPU 계열": "supported_cpu_family", "BIOS 주의": "bios_note",
        "크기": "size_note", "무게": "weight_note", "CPU 지원 URL": "spec_url",
        "판매 상태": "sale_status", "상태 확인일": "status_checked_at",
        "DIMM 슬롯 수": "dimm_slots", "최대 메모리 용량(GB)": "max_memory_gb",
        "최대 메모리 속도(MT/s)": "max_memory_speed_mts", "M.2 슬롯 수": "m2_slots",
        "M.2 PCIe 세대": "m2_pcie_gen", "SATA 포트 수": "sata_ports",
        "CPU 세대별 최소 BIOS 버전": "min_bios", "확장 사양 출처 URL": "expansion_source_url",
    }),
    "RAM": ("ram", "catalog.ram_spec", {
        "메모리 규격": "memory_type", "속도(MT/s)": "speed_mts",
        "총 용량(GB)": "total_capacity_gb", "모듈 구성": "module_config",
        "모듈당 용량(GB)": "capacity_per_module_gb", "CAS 지연시간": "cas_latency",
        "폼팩터": "form_factor", "핀 수": "pin_count", "정격 전압(V)": "voltage_v",
        "ECC": "ecc", "버퍼 방식": "buffer_type", "오버클럭 프로필": "oc_profile",
        "방열판": "heatsink", "RGB": "rgb", "크기": "size_note", "무게": "weight_note",
        "설명서 URL": "spec_url", "비고": "note",
        "방열판 포함 높이(mm)": "height_mm", "높이 출처 URL": "height_source_url",
    }),
    "GPU": ("gpu", "catalog.gpu_spec", {
        "GPU 분류": "gpu_class", "VRAM(GB)": "vram_gb", "메모리 종류": "memory_type",
        "PCIe 인터페이스": "pcie_interface", "소비전력(W)": "power_w",
        "권장 PSU(W)": "recommended_psu_w", "길이(mm)": "length_mm",
        "높이(mm)": "height_mm", "슬롯 두께": "slot_thickness",
        "전원 커넥터": "power_connector", "보조전원": "aux_power", "ECC": "ecc",
        "제품 라인업": "lineup", "사양 기준": "spec_basis", "크기": "size_note",
        "무게": "weight_note", "드라이버 URL": "spec_url",
        "물리 규격 출처 URL": "dimension_source_url", "물리 규격 미확인 사유": "dimension_gap_reason",
        "성능 점수": "perf_score", "성능 점수 출처 URL": "perf_score_source_url",
    }),
    "SSD": ("ssd", "catalog.ssd_spec", {
        "인터페이스": "interface", "프로토콜": "protocol", "폼팩터": "form_factor",
        "방열판": "heatsink", "지원 용량": "capacity_options", "NAND 유형": "nand_type",
        "DRAM": "dram", "PS5 호환": "ps5_compat", "크기": "size_note",
        "무게": "weight_note", "지원 URL": "spec_url", "판매 상태": "sale_status",
        "상태 확인일": "status_checked_at", "비고": "note",
    }),
    "PSU": ("psu", "catalog.psu_spec", {
        "정격 출력(W)": "wattage_w", "효율 등급": "efficiency_rating",
        "폼팩터": "form_factor", "케이블 방식": "cable_type",
        "GPU 전원 커넥터": "gpu_power_connector", "ATX 규격": "atx_spec",
        "모듈러": "modular", "크기": "size_note", "무게": "weight_note",
        "지원 URL": "spec_url", "판매 상태": "sale_status",
        "상태 확인일": "status_checked_at", "비고": "note",
        "길이(mm)": "length_mm", "PCIe 8핀 개수": "pcie_8pin_count",
        "12V-2x6 개수": "connector_12v2x6_count", "길이·커넥터 출처 URL": "dimension_source_url",
    }),
    "Case": ("case", "catalog.case_spec", {
        "지원 메인보드": "supported_motherboard", "CPU 쿨러 높이(mm)": "cpu_cooler_height_mm",
        "GPU 최대 길이(mm)": "gpu_max_length_mm", "PSU 폼팩터": "psu_form_factor",
        "케이스 분류": "case_type", "크기": "size_note", "무게": "weight_note",
        "지원 URL": "spec_url", "판매 상태": "sale_status",
        "상태 확인일": "status_checked_at", "비고": "note",
        "최대 파워 길이(mm)": "max_psu_length_mm", "확장 슬롯 수": "expansion_slots",
        "라디에이터 전면(mm)": "radiator_front_mm", "라디에이터 상단(mm)": "radiator_top_mm",
        "라디에이터 후면(mm)": "radiator_rear_mm", "색상": "color", "확장 사양 출처 URL": "expansion_source_url",
    }),
    "Cooler": ("cooler", "catalog.cooler_spec", {
        "냉각 방식": "cooling_type", "지원 소켓": "supported_socket",
        "쿨러 높이(mm)": "cooler_height_mm", "라디에이터(mm)": "radiator_mm",
        "케이스 확인 항목": "case_check_note", "크기": "size_note", "무게": "weight_note",
        "지원 URL": "spec_url", "판매 상태": "sale_status",
        "상태 확인일": "status_checked_at", "비고": "note",
    }),
}


# 원본에서 숫자로 읽히기 쉽지만 DB 열은 text 인 것(예: 라디에이터 '360' 하나만 있으면 엑셀이 숫자로 준다)
_TEXT_COLUMNS = {"radiator_front_mm", "radiator_top_mm", "radiator_rear_mm", "m2_pcie_gen", "min_bios", "color", "family"}


def _as_text(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value)


def _clean(value):
    """빈 자리표시자('-', 공백)를 None으로. 그 외 값은 그대로 통과."""
    if isinstance(value, str) and value.strip() in ("", "-"):
        return None
    return value


def _upsert_spec(conn, table: str, product_id, row: dict, header_to_col: dict[str, str]) -> None:
    cols = ["product_id"]
    vals = [product_id]
    for header, col in header_to_col.items():
        cols.append(col)
        value = _clean(row.get(header))
        if col in _TEXT_COLUMNS and value is not None and not isinstance(value, str):
            value = _as_text(value)
        vals.append(value)
    placeholders = ", ".join(["%s"] * len(cols))
    update_clause = ", ".join(f"{c} = EXCLUDED.{c}" for c in cols if c != "product_id")
    sql = (
        f"INSERT INTO {table} ({', '.join(cols)}) VALUES ({placeholders}) "
        f"ON CONFLICT (product_id) DO UPDATE SET {update_clause}"
    )
    conn.execute(sql, vals)


def main() -> int:
    if not _XLSX_PATH.exists():
        print(f"파일 없음: {_XLSX_PATH}")
        return 1

    wb = openpyxl.load_workbook(_XLSX_PATH, data_only=True)

    with get_conn() as conn:
        repo = ProductRepo(conn)

        source = repo._one("SELECT id FROM evidence.source WHERE name = %s", (_SOURCE_NAME,))
        if source is None:
            source = repo._one(
                "INSERT INTO evidence.source (name, source_type) VALUES (%s, 'derived') RETURNING id",
                (_SOURCE_NAME,),
            )
        source_id = source["id"]
        merchant_id = repo.upsert_merchant(_MERCHANT_PLATFORM, _MERCHANT_SELLER_ID, _MERCHANT_NAME)

        total = 0
        by_sheet: dict[str, int] = {}
        for sheet_name, (product_type, spec_table, col_map) in _SHEET_SPECS.items():
            if sheet_name not in wb.sheetnames:
                print(f"시트 없음, 건너뜀: {sheet_name}")
                continue
            ws = wb[sheet_name]
            headers = [c.value for c in ws[1]]
            n = 0
            for r in range(2, ws.max_row + 1):
                row = {headers[i]: ws.cell(r, i + 1).value for i in range(len(headers))}
                brand, model = row.get("제조사"), row.get("모델명")
                if not brand or not model:
                    continue

                image_url = _clean(row.get("이미지 URL"))
                product_id = repo.upsert_product(
                    name=f"{brand} {model}", brand=brand, model=model,
                    product_type=product_type, attributes={},
                )
                if image_url:
                    repo._exec(
                        "UPDATE catalog.product SET image_url = %s WHERE id = %s",
                        (image_url, product_id),
                    )
                variant_id = repo.upsert_variant(product_id, "default", attributes={})

                external_offer_id = f"{product_type}:{brand}:{model}".lower().replace(" ", "-")
                purchase_url = _clean(row.get("상품 URL")) or f"https://search.danawa.com/dsearch.php?query={model}"
                offer_id = repo.upsert_offer(variant_id, merchant_id, external_offer_id, purchase_url)

                price = row.get("가격")
                price = float(price) if isinstance(price, (int, float)) else None
                repo.add_observation_if_changed(
                    offer_id, source_id=source_id, observed_at=datetime.now(timezone.utc),
                    price=price, currency="KRW",
                    stock_status="available" if price is not None else "unknown",
                    quality_status="valid" if price is not None else "stale",
                )

                _upsert_spec(conn, spec_table, product_id, row, col_map)
                n += 1
            by_sheet[sheet_name] = n
            total += n
            print(f"  {sheet_name}: {n}개")

        print(f"seed_pc_parts_specs: 총 {total}개 부품 → product/variant/offer/observation + 카테고리 스펙")
    return 0


if __name__ == "__main__":
    sys.exit(main())
